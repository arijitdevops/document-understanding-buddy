"""Per-format document loaders.

Every loader returns ``list[ParsedPage]`` so downstream chunking can attach a
page number to each chunk and citations can point at a page. Formats that have
no natural pagination (TXT, CSV, MD) return a single page, or one page per
slide (PPTX) or per 200 rows (CSV) where that is the more useful unit.

Parsing libraries are synchronous and CPU-bound; :func:`load_document` is async
and runs each parser on a worker thread so the event loop stays responsive.
"""

from __future__ import annotations

import asyncio
import csv
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.services.gemini import GeminiNotConfiguredError, transcribe_image

logger = logging.getLogger(__name__)

#: Pages with at least this many table-ish rows are re-parsed with pdfplumber.
_TABLE_ROW_THRESHOLD = 3
_MAX_CSV_ROWS_PER_PAGE = 200


class UnsupportedDocumentError(ValueError):
    """Raised when no loader can handle the given MIME type / extension."""

    def __init__(self, ext: str, mime: str) -> None:
        super().__init__(
            f"No loader for extension {ext or '(none)'} / MIME {mime or '(none)'}. "
            f"Supported extensions: {', '.join(sorted(supported_extensions()))}"
        )
        self.ext = ext
        self.mime = mime


class DocumentParseError(RuntimeError):
    """Raised when a supported format could not be parsed (corrupt, encrypted...)."""


@dataclass(slots=True)
class ParsedPage:
    """One logical page of extracted text."""

    page_number: int
    text: str
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def is_empty(self) -> bool:
        """``True`` when the page holds no usable text."""
        return not self.text.strip()


def supported_extensions() -> set[str]:
    """Extensions this module can parse."""
    return {
        ".pdf",
        ".docx",
        ".pptx",
        ".csv",
        ".txt",
        ".md",
        ".markdown",
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
    }


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def _looks_table_heavy(text: str) -> bool:
    """Cheap heuristic: several lines with multiple wide gaps suggest a table."""
    rows = 0
    for line in text.splitlines():
        if line.count("  ") >= 2 and len(line.strip()) > 12:
            rows += 1
            if rows >= _TABLE_ROW_THRESHOLD:
                return True
    return False


def _tables_to_text(tables: list[list[list[str | None]]]) -> str:
    """Render pdfplumber tables as pipe-separated rows."""
    lines: list[str] = []
    for table in tables:
        for row in table:
            cells = [(cell or "").replace("\n", " ").strip() for cell in row]
            if any(cells):
                lines.append(" | ".join(cells))
        lines.append("")
    return "\n".join(lines).strip()


def load_pdf(path: Path) -> list[ParsedPage]:
    """Extract text per page with ``pypdf``; re-parse table-heavy pages with ``pdfplumber``.

    Raises:
        DocumentParseError: If the file cannot be opened or is encrypted.
    """
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise DocumentParseError(
            "pypdf is not installed; run pip install -r requirements.txt"
        ) from exc

    try:
        reader = PdfReader(str(path))
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")  # many PDFs use an empty owner password
            except Exception as exc:  # noqa: BLE001
                raise DocumentParseError("The PDF is password protected.") from exc
    except DocumentParseError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"Could not open the PDF: {exc}") from exc

    pages: list[ParsedPage] = []
    table_candidates: list[int] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001 - one bad page must not kill the doc
            logger.warning("pypdf failed on page %d: %s", index, exc)
            text = ""
        if _looks_table_heavy(text):
            table_candidates.append(index)
        pages.append(ParsedPage(page_number=index, text=text, metadata={"parser": "pypdf"}))

    if table_candidates:
        _augment_with_pdfplumber(path, pages, table_candidates)

    if all(page.is_empty for page in pages):
        logger.warning(
            "No text layer found in %s -- it is probably a scan. Upload the pages "
            "as images to use Gemini vision OCR.",
            path.name,
        )
    return pages


def _augment_with_pdfplumber(path: Path, pages: list[ParsedPage], candidates: list[int]) -> None:
    """Append pdfplumber table extraction to the listed pages, in place."""
    try:
        import pdfplumber
    except ImportError:  # pragma: no cover - optional path
        logger.info("pdfplumber not installed; keeping pypdf text for table pages")
        return
    try:
        with pdfplumber.open(str(path)) as pdf:
            for number in candidates:
                if number > len(pdf.pages):
                    continue
                try:
                    tables = pdf.pages[number - 1].extract_tables() or []
                except Exception as exc:  # noqa: BLE001
                    logger.warning("pdfplumber failed on page %d: %s", number, exc)
                    continue
                rendered = _tables_to_text(tables)
                if rendered:
                    page = pages[number - 1]
                    page.text = f"{page.text}\n\n[tables]\n{rendered}".strip()
                    page.metadata["parser"] = "pypdf+pdfplumber"
                    page.metadata["tables"] = len(tables)
    except Exception as exc:  # noqa: BLE001 - enrichment is best-effort
        logger.warning("pdfplumber pass failed for %s: %s", path.name, exc)


# --------------------------------------------------------------------------
# Office formats
# --------------------------------------------------------------------------
def load_docx(path: Path) -> list[ParsedPage]:
    """Extract paragraphs and tables from a .docx file.

    Word has no reliable page concept without rendering, so the document is
    returned as a single page and headings carry the structure instead.
    """
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("python-docx is not installed.") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"Could not open the DOCX: {exc}") from exc

    lines: list[str] = []
    headings = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "") or ""
        if style.lower().startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append(f"{'#' * min(int(level), 6)} {text}")
            headings += 1
        else:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    return [
        ParsedPage(
            page_number=1,
            text="\n\n".join(lines),
            metadata={"parser": "python-docx", "headings": headings},
        )
    ]


def load_pptx(path: Path) -> list[ParsedPage]:
    """Extract text from a .pptx file, one page per slide."""
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError("python-pptx is not installed.") from exc
    try:
        presentation = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"Could not open the PPTX: {exc}") from exc

    pages: list[ParsedPage] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        title = ""
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text_frame.text.strip()
            if not text:
                continue
            if not title:
                title = text.splitlines()[0][:120]
                lines.append(f"# {title}")
                remainder = "\n".join(text.splitlines()[1:]).strip()
                if remainder:
                    lines.append(remainder)
            else:
                lines.append(text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[speaker notes] {notes}")
        pages.append(
            ParsedPage(
                page_number=index,
                text="\n\n".join(lines),
                metadata={"parser": "python-pptx", "slide_title": title},
            )
        )
    return pages


def load_csv(path: Path) -> list[ParsedPage]:
    """Extract a CSV/TSV file, paginating long files so citations stay precise."""
    try:
        with path.open("r", encoding="utf-8-sig", newline="", errors="replace") as handle:
            sample = handle.read(8192)
            handle.seek(0)
            try:
                dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(handle, dialect))
    except OSError as exc:
        raise DocumentParseError(f"Could not read the CSV: {exc}") from exc

    if not rows:
        return [ParsedPage(page_number=1, text="", metadata={"parser": "csv"})]

    header, body = rows[0], rows[1:]
    header_line = " | ".join(cell.strip() for cell in header)
    pages: list[ParsedPage] = []
    for page_index, start in enumerate(
        range(0, max(len(body), 1), _MAX_CSV_ROWS_PER_PAGE), start=1
    ):
        window = body[start : start + _MAX_CSV_ROWS_PER_PAGE]
        lines = [header_line, *(" | ".join(cell.strip() for cell in row) for row in window)]
        pages.append(
            ParsedPage(
                page_number=page_index,
                text="\n".join(lines),
                metadata={"parser": "csv", "row_offset": start, "rows": len(window)},
            )
        )
    return pages


def load_text(path: Path) -> list[ParsedPage]:
    """Read a plain-text or Markdown file as one page."""
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise DocumentParseError(f"Could not read the file: {exc}") from exc
    return [ParsedPage(page_number=1, text=text, metadata={"parser": "text"})]


# --------------------------------------------------------------------------
# Images (Gemini vision OCR)
# --------------------------------------------------------------------------
async def load_image(path: Path, mime: str) -> list[ParsedPage]:
    """Transcribe an image with Gemini vision (through LangChain).

    Raises:
        GeminiNotConfiguredError: If no API key is configured.
        DocumentParseError: If transcription failed.
    """
    try:
        data = await asyncio.to_thread(path.read_bytes)
        text = await transcribe_image(data, mime)
    except GeminiNotConfiguredError:
        raise
    except Exception as exc:  # noqa: BLE001 - surfaced as a parse failure
        raise DocumentParseError(f"Image transcription failed: {exc}") from exc
    return [
        ParsedPage(
            page_number=1,
            text=text,
            metadata={"parser": "gemini-vision", "mime": mime},
        )
    ]


# --------------------------------------------------------------------------
# Dispatch
# --------------------------------------------------------------------------
_SYNC_LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".pptx": load_pptx,
    ".csv": load_csv,
    ".txt": load_text,
    ".md": load_text,
    ".markdown": load_text,
}

_MIME_TO_EXT = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    "text/csv": ".csv",
    "text/plain": ".txt",
    "text/markdown": ".md",
}

_IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}
_IMAGE_MIME = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}


async def load_document(path: Path | str, mime: str = "") -> list[ParsedPage]:
    """Parse ``path`` into pages, dispatching on MIME type *and* extension.

    The extension wins when both are usable, because browsers frequently send
    ``application/octet-stream``; the MIME type is the fallback for files that
    arrive without a meaningful suffix.

    Raises:
        UnsupportedDocumentError: If neither signal maps to a loader.
        DocumentParseError: If the parser rejected the file.
    """
    path = Path(path)
    ext = path.suffix.lower()
    mime_norm = (mime or "").split(";")[0].strip().lower()

    if ext not in supported_extensions():
        mapped = _MIME_TO_EXT.get(mime_norm)
        if mapped is None and mime_norm.startswith("image/"):
            mapped = ".png"
        if mapped is None:
            raise UnsupportedDocumentError(ext, mime_norm)
        logger.info("Extension %r unknown; falling back to MIME %s", ext, mime_norm)
        ext = mapped

    if ext in _IMAGE_EXTS:
        return await load_image(
            path, mime_norm if mime_norm.startswith("image/") else _IMAGE_MIME.get(ext, "image/png")
        )

    loader = _SYNC_LOADERS.get(ext)
    if loader is None:  # pragma: no cover - table above is exhaustive
        raise UnsupportedDocumentError(ext, mime_norm)

    logger.info("Parsing %s with %s", path.name, loader.__name__)
    pages = await asyncio.to_thread(loader, path)
    return [page for page in pages if not page.is_empty] or pages[:1]
